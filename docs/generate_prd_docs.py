import os
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx(filename):
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x2D, 0x37, 0x48) # Dark Slate

    # --- TITLE ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("PRODUCT REQUIREMENT DOCUMENT (PRD)")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x06, 0x5F, 0x46) # Emerald Dark

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Implementasi Sertifikasi Carbon Offset Berbasis Blockchain (Polygon)\npada Platform Web CAMAR (Tugas Akhir / Skripsi)")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # --- METADATA TABLE ---
    table_meta = doc.add_table(rows=6, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False
    meta_data = [
        ("Nama Proyek", "CAMAR (Carbon Offset Platform)"),
        ("Modul / Fitur", "Sertifikasi Digital Terdesentralisasi & Verifikasi On-Chain"),
        ("Teknologi Utama", "PHP 8.2+ / Laravel 12, MySQL, Node.js (ethers.js), Solidity (Smart Contract)"),
        ("Jaringan Blockchain", "Polygon PoS (Amoy Testnet) - Bebas Gas Fee Nyata"),
        ("Standar Token", "Soulbound Token (ERC-721 Non-Transferable) & SHA-256 Hash Notarization"),
        ("Status Dokumen", "Versi 1.0 (Dokumen Kebutuhan Resmi untuk Tugas Akhir)")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = table_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F0FDF4") # Light Emerald
        set_cell_background(c1, "FAFAFA")
        set_cell_margins(c0, 100, 100, 150, 150)
        set_cell_margins(c1, 100, 100, 150, 150)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
        return h

    # --- SECTION 1 ---
    add_h1("1. Latar Belakang & Rumusan Masalah")
    
    add_h2("1.1 Konteks Proyek CAMAR")
    p = doc.add_paragraph(
        "Platform CAMAR (Carbon Offset Platform) adalah sistem e-marketplace berbasis Laravel yang mempertemukan penyedia proyek penyerapan karbon (Seller) dengan pembeli (Buyer/Perusahaan) yang hendak mengompensasi emisi gas rumah kaca mereka. Saat ini sistem telah memiliki fitur penghitungan jejak emisi, pemilihan proyek, pembayaran via Midtrans, hingga verifikasi transaksi oleh Auditor Pemerintah dan penerbitan nomor sertifikat internal (CAMAR-CERT-YYYYMMDD-XXXXXX)."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    add_h2("1.2 Permasalahan yang Dihadapi")
    bullet1 = doc.add_paragraph(style='List Bullet')
    r1 = bullet1.add_run("Kerapuhan Basis Data Terpusat (Centralized DB Vulnerability): ")
    r1.font.bold = True
    bullet1.add_run("Pencatatan sertifikat pada database MySQL tradisional rentan terhadap manipulasi sepihak oleh administrator yang memiliki akses root database.")

    bullet2 = doc.add_paragraph(style='List Bullet')
    r2 = bullet2.add_run("Isu Klaim Ganda (Double Counting): ")
    r2.font.bold = True
    bullet2.add_run("Sertifikat PDF tanpa verifikasi kriptografis terdistribusi dapat digunakan atau diklaim berulang kali oleh pihak pembeli pada audit lingkungan yang berbeda.")

    bullet3 = doc.add_paragraph(style='List Bullet')
    r3 = bullet3.add_run("Ketiadaan Mekanisme Verifikasi Publik: ")
    r3.font.bold = True
    bullet3.add_run("Pihak ketiga (auditor eksternal, regulator, masyarakat) tidak dapat menguji keaslian sertifikat tanpa harus memiliki hak akses login ke dalam sistem CAMAR.")

    add_h2("1.3 Solusi yang Diajukan")
    p = doc.add_paragraph(
        "Mengembangkan lapisan otentikasi berbasis Smart Contract pada blockchain Polygon Amoy Testnet. Setiap sertifikat yang disahkan oleh Auditor Pemerintah akan dihitung nilai sidik jari digitalnya (SHA-256) dan dicetak sebagai Soulbound Token (SBT) yang kekal (immutable) serta tidak dapat dipindahtangankan, dilengkapi QR Code verifikasi independen."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(10)

    # --- SECTION 2 ---
    add_h1("2. Arsitektur Sistem (Opsi C: Hybrid Worker)")
    p = doc.add_paragraph(
        "Sistem menggunakan arsitektur Hybrid Microservice di mana Laravel berperan sebagai pengelola logika bisnis dan database relasional, sementara sebuah skrip mandiri Node.js bertindak sebagai jembatan (Worker Bridge) ke jaringan Blockchain melalui pustaka ethers.js."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # Diagram Table
    tbl_diag = doc.add_table(rows=5, cols=2)
    tbl_diag.alignment = WD_TABLE_ALIGNMENT.CENTER
    diag_steps = [
        ("Langkah 1: Approval & Hashing", "Auditor menyetujui transaksi di Admin Panel. Laravel men-generate nomor seri sertifikat dan menghitung nilai hash SHA-256 dari parameter dokumen sertifikat."),
        ("Langkah 2: Eksekusi Worker", "Laravel memicu worker Node.js (via Process::run atau background queue) dengan membawa payload data sertifikat dan hash."),
        ("Langkah 3: Transaksi On-Chain", "Worker menandatangani transaksi menggunakan Private Key sistem (Gasless bagi user) dan memanggil fungsi issueCertificate() pada Smart Contract Polygon Amoy."),
        ("Langkah 4: Penyimpanan TxHash", "Jaringan mengonfirmasi blok dan mengembalikan Transaction Hash (TxHash) & Token ID. Laravel menyimpan data ini ke database orders."),
        ("Langkah 5: Tampilan & Verifikasi", "Buyer menerima sertifikat berstempel on-chain dengan QR Code yang dapat dipindai secara publik menuju Block Explorer (Polygonscan).")
    ]
    for i, (step, desc) in enumerate(diag_steps):
        r = tbl_diag.rows[i]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "ECFDF5")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 80, 80, 120, 120)
        set_cell_margins(c1, 80, 80, 120, 120)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(step)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(desc)
        r1.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 3 ---
    add_h1("3. Spesifikasi Perubahan Basis Data (Database Schema)")
    p = doc.add_paragraph("Tabel orders akan diperluas dengan kolom metadata blockchain berikut:")
    p.paragraph_format.space_after = Pt(6)

    tbl_db = doc.add_table(rows=8, cols=3)
    tbl_db.alignment = WD_TABLE_ALIGNMENT.CENTER
    db_fields = [
        ("Nama Kolom", "Tipe Data", "Fungsi & Deskripsi"),
        ("certificate_hash", "VARCHAR(64)", "Nilai SHA-256 hash unik dokumen untuk menjamin integritas data."),
        ("blockchain_network", "VARCHAR(50)", "Identitas jaringan (Default: Polygon Amoy Testnet)."),
        ("contract_address", "VARCHAR(42)", "Alamat Smart Contract sertifikat yang telah dideploy."),
        ("blockchain_tx_hash", "VARCHAR(66)", "Hash transaksi bukti pencatatan on-chain di Polygon."),
        ("blockchain_token_id", "BIGINT UNSIGNED", "Nomor urut token sertifikat di dalam smart contract."),
        ("blockchain_status", "ENUM", "Status pencatatan: unminted, pending, minted, failed."),
        ("blockchain_minted_at", "TIMESTAMP", "Waktu persis konfirmasi blok transaksi blockchain.")
    ]
    for i, (col, dtype, desc) in enumerate(db_fields):
        r = tbl_db.rows[i]
        c0, c1, c2 = r.cells[0], r.cells[1], r.cells[2]
        c0.width = Inches(1.8)
        c1.width = Inches(1.5)
        c2.width = Inches(3.2)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(col)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(dtype)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(desc)

        if i == 0:
            for c, p, r in [(c0, p0, r0), (c1, p1, r1), (c2, p2, r2)]:
                set_cell_background(c, "065F46")
                set_cell_margins(c, 100, 100, 100, 100)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9.5)
        else:
            set_cell_background(c0, "F8FAFC")
            set_cell_background(c1, "F8FAFC")
            set_cell_background(c2, "FFFFFF")
            for c in (c0, c1, c2):
                set_cell_margins(c, 70, 70, 100, 100)
            r0.font.bold = True
            r0.font.size = Pt(9)
            r1.font.size = Pt(9)
            r2.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 4 ---
    add_h1("4. Kebutuhan Fungsional (Functional Requirements)")
    
    fr_items = [
        ("FR-1: Smart Contract Soulbound (Solidity)", [
            "Kontrak pintar diberi nama CamarCarbonCertificate.sol berbasis standar ERC-721.",
            "Fungsi issueCertificate() mencatat: nomor sertifikat, penerima, tonase CO2, hash dokumen, dan timestamp.",
            "Fitur Soulbound: Fungsi transfer (transferFrom / safeTransferFrom) diblokir untuk mencegah penjualan kembali atas emisi yang telah dikonsumsi (retired).",
            "Fungsi verifikasi publik verifyCertificateHash(certNumber, certHash) yang dapat dieksekusi tanpa gas fee."
        ]),
        ("FR-2: Node.js Worker Bridge (ethers.js)", [
            "Diletakkan di direktori blockchain/scripts/issue-certificate.js.",
            "Menggunakan RPC provider Polygon Amoy dan menandatangani transaksi menggunakan sistem relayer wallet.",
            "Menghasilkan output JSON terstruktur yang berisi status transaksi, tx_hash, dan token_id."
        ]),
        ("FR-3: Integrasi Pengesahan di Laravel", [
            "Dipicu saat Auditor Pemerintah mengubah status transaksi menjadi completed di TransactionManagementController.",
            "Menghasilkan SHA-256 hash secara deterministik dari atribut transaksi.",
            "Mencatat seluruh jejak audit ke dalam tabel admin_activity_logs."
        ]),
        ("FR-4: Template Sertifikat Interaktif (certificate.blade.php)", [
            "Menampilkan Digital Seal 'Verified on Polygon Blockchain'.",
            "Menampilkan rincian teknis: Network, TxHash, dan Document Hash.",
            "Menyematkan Dynamic QR Code yang mengarah ke URL verifikasi publik."
        ]),
        ("FR-5: Portal Verifikasi Publik (/verify/certificate/{code})", [
            "Dapat diakses oleh siapa saja tanpa autentikasi login.",
            "Menampilkan status keaslian data, waktu blok, dan tautan Polygonscan Explorer.",
            "Mampu mendeteksi manipulasi data: Apabila data di database lokal diubah, sistem mendeteksi ketidakcocokan nilai hash dengan blockchain."
        ])
    ]

    for title, bullets in fr_items:
        add_h2(title)
        for b in bullets:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(2)
            bp.paragraph_format.line_spacing = 1.1
            r = bp.add_run(b)
            r.font.size = Pt(10)

    # --- SECTION 5 ---
    add_h1("5. Skenario Pengujian untuk Laporan Skripsi (Bab 4)")
    tests = [
        ("1. Pengujian Fungsional (Black-Box Testing)", "Menguji seluruh siklus transaksi mulai dari kalkulasi emisi, pembayaran Midtrans, approval auditor, eksekusi smart contract, hingga penerbitan sertifikat."),
        ("2. Pengujian Integritas Data (Data Tampering Test)", "Melakukan pengujian ketahanan dengan sengaja mengubah kuota tonase atau nama pembeli di tabel MySQL. Hasil yang diharapkan: Halaman verifikasi publik berhasil mendeteksi anomali 'Data Telah Dimanipulasi' karena hash lokal tidak cocok dengan hash on-chain."),
        ("3. Pengujian Kinerja & Gas Fee (Performance Benchmark)", "Mengukur waktu latensi penandatanganan transaksi (signing), konfirmasi blok di jaringan Polygon Amoy (rata-rata 2–4 detik), serta konsumsi gas fee per transaksi.")
    ]
    for t_title, t_desc in tests:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        rt = p.add_run(f"• {t_title}: ")
        rt.font.bold = True
        p.add_run(t_desc)

    # --- SECTION 6 ---
    add_h1("6. Jadwal & Milestone Pengerjaan")
    p = doc.add_paragraph(
        "Implementasi dibagi menjadi 4 tahap: (1) Setup Smart Contract & Deploy ke Polygon Amoy, (2) Pembuatan Worker ethers.js & Migration Database, (3) Modifikasi Controller & Template Sertifikat Blade, (4) Pembuatan Halaman Verifikasi Publik & Pengujian Integritas."
    )
    p.paragraph_format.space_after = Pt(15)

    doc.save(filename)
    print(f"DOCX created: {filename}")

def create_html(filename):
    html_content = """<!DOCTYPE html>
<html lang="id">
<head>
    <meta charset="UTF-8">
    <title>PRD - Sistem Sertifikasi Carbon Offset Berbasis Blockchain (CAMAR)</title>
    <style>
        @page {
            size: A4;
            margin: 18mm 16mm 18mm 16mm;
            @bottom-right {
                content: counter(page);
                font-family: Arial, sans-serif;
                font-size: 9pt;
                color: #64748b;
            }
        }
        body {
            font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, Helvetica, Arial, sans-serif;
            color: #1e293b;
            background: #ffffff;
            line-height: 1.5;
            margin: 0;
            padding: 0;
            font-size: 10pt;
        }
        .header-box {
            text-align: center;
            border-bottom: 3px double #059669;
            padding-bottom: 16px;
            margin-bottom: 20px;
        }
        .title {
            color: #065f46;
            font-size: 20pt;
            font-weight: 800;
            margin: 0 0 6px 0;
            letter-spacing: 0.5px;
        }
        .subtitle {
            color: #475569;
            font-size: 11pt;
            font-weight: 600;
            margin: 0;
        }
        .meta-table {
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 22px;
            font-size: 9pt;
        }
        .meta-table td {
            padding: 7px 12px;
            border: 1px solid #e2e8f0;
        }
        .meta-table td.label {
            background-color: #f0fdf4;
            color: #065f46;
            font-weight: 700;
            width: 28%;
        }
        .meta-table td.value {
            background-color: #f8fafc;
            color: #1e293b;
        }
        h2 {
            color: #065f46;
            font-size: 13pt;
            border-left: 4px solid #059669;
            padding-left: 10px;
            margin-top: 20px;
            margin-bottom: 8px;
            page-break-after: avoid;
        }
        h3 {
            color: #0f766e;
            font-size: 10.5pt;
            margin-top: 12px;
            margin-bottom: 6px;
            page-break-after: avoid;
        }
        p {
            margin: 0 0 8px 0;
            text-align: justify;
        }
        ul {
            margin: 4px 0 10px 20px;
            padding: 0;
        }
        li {
            margin-bottom: 4px;
        }
        .data-table {
            width: 100%;
            border-collapse: collapse;
            margin: 10px 0 16px 0;
            font-size: 8.5pt;
        }
        .data-table th {
            background-color: #065f46;
            color: #ffffff;
            font-weight: 700;
            padding: 8px 10px;
            text-align: left;
            border: 1px solid #065f46;
        }
        .data-table td {
            padding: 6px 10px;
            border: 1px solid #cbd5e1;
            vertical-align: top;
        }
        .data-table tr:nth-child(even) td {
            background-color: #f8fafc;
        }
        .badge {
            display: inline-block;
            background: #d1fae5;
            color: #065f46;
            font-weight: bold;
            padding: 2px 6px;
            border-radius: 4px;
            font-size: 8pt;
        }
        .callout {
            background: #f0fdf4;
            border: 1px solid #a7f3d0;
            border-left: 4px solid #059669;
            padding: 10px 14px;
            margin: 12px 0;
            border-radius: 0 6px 6px 0;
            font-size: 9.5pt;
        }
        .page-break {
            page-break-before: always;
        }
    </style>
</head>
<body>

    <div class="header-box">
        <h1 class="title">PRODUCT REQUIREMENT DOCUMENT (PRD)</h1>
        <div class="subtitle">Implementasi Sertifikasi Carbon Offset Berbasis Blockchain (Polygon Amoy Testnet)<br>pada Platform Web CAMAR (Tugas Akhir / Skripsi)</div>
    </div>

    <table class="meta-table">
        <tr>
            <td class="label">Nama Proyek</td>
            <td class="value"><strong>CAMAR (Carbon Offset Platform)</strong></td>
        </tr>
        <tr>
            <td class="label">Fitur Utama</td>
            <td class="value">Sertifikasi Digital Terdesentralisasi & Verifikasi Keaslian On-Chain</td>
        </tr>
        <tr>
            <td class="label">Arsitektur Teknis</td>
            <td class="value">Laravel 12 (PHP 8.2+), MySQL, Node.js + ethers.js (Worker Bridge), Solidity</td>
        </tr>
        <tr>
            <td class="label">Jaringan Blockchain</td>
            <td class="value">Polygon PoS (Amoy Testnet) - Ramah Lingkungan & Bebas Biaya Transaksi Riil</td>
        </tr>
        <tr>
            <td class="label">Standar Smart Contract</td>
            <td class="value">Soulbound Token (ERC-721 Non-Transferable) & SHA-256 Hash Notarization</td>
        </tr>
        <tr>
            <td class="label">Target Evaluasi</td>
            <td class="value">Tugas Akhir / Skripsi Program Studi Informatika / Sistem Informasi / RPL</td>
        </tr>
    </table>

    <h2>1. Latar Belakang & Rumusan Masalah</h2>
    <h3>1.1 Konteks Proyek CAMAR</h3>
    <p>
        Platform CAMAR dirancang untuk menjembatani penyedia proyek penyerapan karbon (Seller) dengan pembeli emisi (Buyer). Saat ini, CAMAR telah dilengkapi alur pemesanan lengkap: penghitungan emisi via kalkulator, checkout via Midtrans payment gateway, verifikasi transaksi oleh Auditor Pemerintah, hingga penerbitan nomor sertifikat internal (<code>CAMAR-CERT-YYYYMMDD-XXXXXX</code>) yang dapat dicetak oleh Buyer.
    </p>

    <h3>1.2 Permasalahan yang Dihadapi (Problem Statement)</h3>
    <ul>
        <li><strong>Kerapuhan Basis Data Terpusat (Centralized DB Vulnerability):</strong> Penyimpanan sertifikat pada database MySQL tradisional memiliki risiko manipulasi sepihak oleh pengguna dengan hak akses tingkat tinggi.</li>
        <li><strong>Isu Klaim Ganda (Double Counting):</strong> Sertifikat konvensional berformat PDF tanpa bukti otentikasi on-chain rentan disalahgunakan atau diklaim berulang kali pada audit kepatuhan lingkungan yang berbeda.</li>
        <li><strong>Ketiadaan Verifikasi Publik Mandiri:</strong> Publik, akademisi, atau auditor eksternal tidak memiliki akses independen untuk mengonfirmasi keaslian sertifikat tanpa mendaftar ke platform CAMAR.</li>
    </ul>

    <h3>1.3 Solusi yang Diajukan</h3>
    <p>
        Membangun lapisan verifikasi berbasis Smart Contract di jaringan <strong>Polygon Amoy Testnet</strong>. Setiap sertifikat yang diterbitkan akan dihitung sidik jari digitalnya (SHA-256) dan dicetak sebagai <strong>Soulbound Token (SBT)</strong>. Sertifikat dilengkapi QR Code dinamis menuju halaman verifikasi publik yang tersinkronisasi langsung dengan <em>Block Explorer</em> (Polygonscan).
    </p>

    <h2>2. Arsitektur Sistem (Opsi C: Hybrid Worker)</h2>
    <div class="callout">
        <strong>Pola Arsitektur:</strong> Laravel bertindak sebagai pengatur alur transaksi, antrean, dan database lokal, sedangkan worker mandiri berbasis <code>ethers.js</code> mengeksekusi penandatanganan transaksi dan interaksi smart contract secara asinkron (Gasless bagi Buyer).
    </div>

    <table class="data-table">
        <thead>
            <tr>
                <th style="width: 25%;">Tahapan Alur</th>
                <th>Mekanisme Kerja & Integrasi</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td><strong>1. Approval & Hashing</strong></td>
                <td>Auditor menyetujui transaksi di Admin Panel. Laravel men-generate nomor sertifikat dan menghitung SHA-256 hash dari data kanonikal sertifikat.</td>
            </tr>
            <tr>
                <td><strong>2. Panggilan Worker</strong></td>
                <td>Laravel memicu skrip <code>issue-certificate.js</code> (via <code>Process::run</code> atau Queue Job) dengan membawa data pesanan dan hash dokumen.</td>
            </tr>
            <tr>
                <td><strong>3. Eksekusi On-Chain</strong></td>
                <td>Worker menandatangani transaksi menggunakan Private Key sistem dan memanggil fungsi <code>issueCertificate()</code> pada Smart Contract Polygon Amoy.</td>
            </tr>
            <tr>
                <td><strong>4. Konfirmasi & Penyimpanan</strong></td>
                <td>Jaringan mengembalikan Transaction Hash (TxHash) & Token ID. Laravel memperbarui record transaksi di tabel <code>orders</code>.</td>
            </tr>
            <tr>
                <td><strong>5. Verifikasi Publik</strong></td>
                <td>Buyer menerima sertifikat berstempel on-chain dengan QR Code dinamis yang langsung dapat dipindai oleh pihak luar.</td>
            </tr>
        </tbody>
    </table>

    <div class="page-break"></div>

    <h2>3. Spesifikasi Basis Data (Database Schema)</h2>
    <p>Tabel <code>orders</code> ditambahkan kolom metadata blockchain untuk menjamin integritas dan rekam jejak:</p>
    <table class="data-table">
        <thead>
            <tr>
                <th style="width: 25%;">Nama Kolom</th>
                <th style="width: 20%;">Tipe Data</th>
                <th>Fungsi & Deskripsi</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td><code>certificate_hash</code></td>
                <td>VARCHAR(64), Nullable</td>
                <td>Sidik jari digital dokumen (SHA-256) untuk validasi integritas anti-manipulasi.</td>
            </tr>
            <tr>
                <td><code>blockchain_network</code></td>
                <td>VARCHAR(50), Nullable</td>
                <td>Nama jaringan blockchain (Default: <em>Polygon Amoy Testnet</em>).</td>
            </tr>
            <tr>
                <td><code>contract_address</code></td>
                <td>VARCHAR(42), Nullable</td>
                <td>Alamat kontrak pintar (Smart Contract) yang digunakan pada saat penerbitan.</td>
            </tr>
            <tr>
                <td><code>blockchain_tx_hash</code></td>
                <td>VARCHAR(66), Nullable</td>
                <td>Hash transaksi unik dari jaringan Polygon sebagai bukti on-chain permanen.</td>
            </tr>
            <tr>
                <td><code>blockchain_token_id</code></td>
                <td>BIGINT UNSIGNED, Nullable</td>
                <td>ID token unik on-chain (Nomor urut NFT/Soulbound Token).</td>
            </tr>
            <tr>
                <td><code>blockchain_status</code></td>
                <td>ENUM('unminted','pending','minted','failed')</td>
                <td>Status sinkronisasi blockchain (Default: <code>unminted</code>).</td>
            </tr>
            <tr>
                <td><code>blockchain_minted_at</code></td>
                <td>TIMESTAMP, Nullable</td>
                <td>Waktu persis konfirmasi blok transaksi di jaringan blockchain.</td>
            </tr>
        </tbody>
    </table>

    <h2>4. Rincian Kebutuhan Fungsional (Functional Requirements)</h2>
    <ul>
        <li><strong>FR-1 (Smart Contract Soulbound):</strong>
            <ul>
                <li>Ditulis dalam bahasa Solidity (<code>CamarCarbonCertificate.sol</code>) berbasis standar ERC-721.</li>
                <li>Fungsi <code>issueCertificate()</code> mencatat: nomor sertifikat, alamat penerima, tonase CO2, hash dokumen, dan waktu penerbitan.</li>
                <li>Fungsi transfer dinonaktifkan (Soulbound) sehingga token tidak dapat diperjualbelikan kembali di secondary market.</li>
            </ul>
        </li>
        <li><strong>FR-2 (Node.js Worker Bridge):</strong>
            <ul>
                <li>Berada pada direktori <code>blockchain/scripts/issue-certificate.js</code>.</li>
                <li>Menghubungkan aplikasi ke RPC Node Polygon Amoy menggunakan pustaka <code>ethers.js</code> v6.</li>
                <li>Mengembalikan keluaran terstruktur berformat JSON (status, tx_hash, token_id, gas_used).</li>
            </ul>
        </li>
        <li><strong>FR-3 (Modul Penerbitan di Laravel):</strong>
            <ul>
                <li>Terintegrasi pada aksi <code>TransactionManagementController@update</code> saat status diubah menjadi <code>completed</code>.</li>
                <li>Mencatat seluruh rekam jejak penerbitan ke tabel <code>admin_activity_logs</code>.</li>
            </ul>
        </li>
        <li><strong>FR-4 (Pembaruan Template Sertifikat Blade):</strong>
            <ul>
                <li>Memperkaya tampilan <code>resources/views/main_page/dashboard-buyer/certificate.blade.php</code> dengan badge <em>"Verified on Polygon Blockchain"</em>.</li>
                <li>Menampilkan TxHash, SHA-256 Hash, serta Dynamic QR Code verifikasi publik.</li>
            </ul>
        </li>
        <li><strong>FR-5 (Portal Verifikasi Publik):</strong>
            <ul>
                <li>Endpoint publik <code>/verify/certificate/{code}</code> yang dapat diakses tanpa autentikasi.</li>
                <li>Mendeteksi manipulasi data lokal (Tamper-evident system).</li>
            </ul>
        </li>
    </ul>

    <h2>5. Skenario Pengujian untuk Laporan Skripsi (Bab 4)</h2>
    <div class="callout">
        <strong>Materi Pengujian Akademis:</strong>
        <ol style="margin: 4px 0 0 16px; padding: 0;">
            <li><strong>Pengujian Fungsional (Black-box):</strong> Memvalidasi seluruh skenario operasional dari pembayaran hingga sertifikat berhasil di-mint.</li>
            <li><strong>Pengujian Integritas Dokumen (Tamper-evident Test):</strong> Memanipulasi nilai data transaksi langsung pada database MySQL, lalu membuktikan bahwa portal verifikasi berhasil mendeteksi anomali ketidakcocokan hash dengan blockchain.</li>
            <li><strong>Pengujian Kinerja & Gas Consumption:</strong> Mengukur waktu respon signing transaksi, durasi konfirmasi blok Polygon (2–4 detik), dan kebutuhan gas fee.</li>
        </ol>
    </div>

</body>
</html>"""
    with open(filename, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"HTML created: {filename}")

if __name__ == '__main__':
    base_dir = r"c:\xampp\htdocs\Carbon Marketplace"
    docs_dir = os.path.join(base_dir, "camar", "docs")
    os.makedirs(docs_dir, exist_ok=True)

    docx_path1 = os.path.join(base_dir, "PRD_Blockchain_Carbon_Certificate_CAMAR.docx")
    docx_path2 = os.path.join(docs_dir, "PRD_Blockchain_Carbon_Certificate_CAMAR.docx")
    html_path = os.path.join(docs_dir, "PRD_Blockchain_Carbon_Certificate_CAMAR.html")

    create_docx(docx_path1)
    create_docx(docx_path2)
    create_html(html_path)
